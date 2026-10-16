import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook("glass_citations_report.xlsx")

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
regular_font = Font(name="Calibri", size=10)
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

for ws in wb.worksheets:
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"
    
    # Format Headers
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Format Data Rows & Set Column Widths
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = regular_font
            cell.border = thin_border
            if cell.column_letter in ['B', 'E', 'G']:  # Platform, Category, Date
                cell.alignment = Alignment(horizontal="center", vertical="top")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top")
                
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if cell.row == 1:
                max_len = max(max_len, len(val) + 4)
            else:
                max_len = max(max_len, min(len(val), 60))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

wb.save("glass_citations_report.xlsx")
print("Formatted glass_citations_report.xlsx successfully!")
