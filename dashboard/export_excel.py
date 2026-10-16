"""Export a dashboard project to a plain Excel workbook.

    python dashboard/export_excel.py education --out edu_ai_visibility.xlsx \
        --previous "ChatGPT run 1 (shared chat)=edu_output_before_fix.csv" \
        --previous "ChatGPT run 2 (shared chat)=edu_output_shared_chat.csv"

Metrics come from the project's current source CSV, computed the same way as
the dashboard. --previous adds older ChatGPT runs to the source sheets and raw data only.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

import data

CELL_LIMIT = 32_000
PLATFORM_LABEL = {"chatgpt": "ChatGPT", "chatgpt_anon": "ChatGPT (logged out)", "gemini": "Gemini"}
URL_TYPE_RULES = [
    (r"(^|\.)wikipedia\.org$", "Wikipedia"),
    (r"(shiksha|collegedunia|careers360|collegedekho|getmyuni|nirfindia|jagranjosh|universitykart|collegesearch)\.", "Education portal"),
    (r"(timesofindia|hindustantimes|ndtv|indianexpress|economictimes|livemint|business-standard|moneycontrol|news18|"
     r"indiatoday|thehindu|firstpost|scroll|thewire|theprint|outlookindia|deccanherald|financialexpress|dnaindia|"
     r"zeenews|abplive|republicworld|wionews|thequint|newslaundry|reuters|bloomberg|forbes|fortune|bbc|cnn|nytimes|"
     r"theguardian|washingtonpost|wsj|ft\.com|apnews|aljazeera|cnbc|businessinsider|time\.com|vogue|people\.com|"
     r"independent\.co\.uk|telegraph|dailymail|nypost)", "News / media"),
    (r"\.gov(\.in)?$|\.nic\.in$|sci\.gov\.in", "Government"),
    (r"(^|\.)(youtube\.com|youtu\.be)$", "Video"),
    (r"(^|\.)(reddit\.com|quora\.com|facebook\.com|instagram\.com|linkedin\.com|x\.com|twitter\.com|threads\.net)$", "Social Media"),
    (r"\.edu(\.in)?$|\.ac\.in$|\.ac\.uk$|university|college|institute", "University / college"),
]


def url_type(url: str) -> str:
    domain = data.domain_of(url)
    for pattern, label in URL_TYPE_RULES:
        if re.search(pattern, domain, re.I):
            return label
    return "Website"


def platform_name(p: str) -> str:
    return PLATFORM_LABEL.get(p, p)


def clip(text: str) -> str:
    return text if len(text) <= CELL_LIMIT else text[:CELL_LIMIT] + " …[truncated]"


def payload_streams(raw: str, prompt: str) -> str:
    """The answer stream(s) in a raw_payload that belong to this prompt."""
    if not raw:
        return ""
    try:
        bodies = [b for b in json.loads(raw) if isinstance(b, str)]
    except ValueError:
        return raw
    needles = {prompt, json.dumps(prompt)[1:-1], json.dumps(prompt, ensure_ascii=False)[1:-1]}
    needles |= {n.replace("&", "\\u0026") for n in needles}
    own = [b for b in bodies if any(n in b for n in needles)]
    return "\n".join(own)


def read_run(path: Path, platforms: set[str] | None = None) -> list[dict]:
    """One record per answer: prompt, platform, date, text, cited, read, searches."""
    answers: dict[tuple, dict] = {}
    last: dict[tuple, tuple] = {}
    with path.open(encoding="utf-8-sig", newline="") as fh:
        for row in csv.DictReader(fh):
            prompt = (row.get("prompt") or "").strip()
            platform = (row.get("platform") or "").strip()
            if not prompt or (platforms and platform not in platforms):
                continue
            date = (row.get("response_date") or "")[:10]
            text = row.get("response_content") or ""
            key = (prompt, platform, date)
            akey = last[key] if (not text.strip() and key in last) else (prompt, platform, date, text)
            last[key] = akey
            ans = answers.setdefault(akey, {"prompt": prompt, "platform": platform, "date": date, "text": text,
                                            "cited": [], "read": [], "queries": []})
            url = (row.get("url") or "").strip()
            if url and not data.is_widget(url):
                clean = data.clean_url(url)
                if clean not in ans["cited"]:
                    ans["cited"].append(clean)
            if row.get("raw_payload") and not ans["read"] and not ans["queries"]:
                if platform.startswith("gemini"):
                    ans["queries"], ans["read"] = data.mine_payload(row["raw_payload"])
                    for u in data.payload_gemini_citations(row["raw_payload"]):
                        if u not in ans["cited"]:
                            ans["cited"].append(u)
                else:
                    ans["queries"], ans["read"] = data.mine_payload(
                        json.dumps([payload_streams(row["raw_payload"], prompt)]))
    for ans in answers.values():
        cited = {data.url_key(u) for u in ans["cited"]}
        ans["read_only"] = [u for u in ans["read"] if data.url_key(u) not in cited]
    return list(answers.values())


def write_sheet(wb: Workbook, title: str, header: list[str], rows: list[list], widths: dict[int, int] | None = None,
                pct_cols: tuple[int, ...] = (), wrap_cols: tuple[int, ...] = (), intro: list[str] | None = None):
    ws = wb.create_sheet(title)
    r = 1
    for line in intro or []:
        ws.cell(row=r, column=1, value=line)
        r += 1
    if intro:
        r += 1
    for c, h in enumerate(header, start=1):
        ws.cell(row=r, column=c, value=h).font = Font(bold=True)
    header_row = r
    for values in rows:
        r += 1
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=clip(v) if isinstance(v, str) else v)
            if c in pct_cols and isinstance(v, (int, float)):
                cell.number_format = "0%"
            if c in wrap_cols:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                cell.alignment = Alignment(vertical="top")
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(header))}{max(r, header_row)}"
    for c in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(c, 16)
    return ws


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("project")
    ap.add_argument("--out", required=True)
    ap.add_argument("--previous", action="append", default=[], help='"Label=path.csv" older ChatGPT runs')
    args = ap.parse_args()

    project = next(p for p in data.load_config()["projects"] if p["id"] == args.project)
    brands = list(project["brands"])
    specs = [data.brand_spec(v) for v in project["brands"].values()]
    patterns = [(i, re.compile(m, re.I)) for i, (m, _) in enumerate(specs)]
    owned = [(site, i) for i, (_, sites) in enumerate(specs) for site in sites]
    own_brand = project.get("own_brand") or brands[0]
    own = brands.index(own_brand)

    def owner(url: str) -> str:
        b = data.owner_of(data.domain_of(url), owned)
        return brands[b] if b is not None else ""

    current = [a for path in data.source_files(project) for a in read_run(path)]
    for a in current:
        a["brands"] = data.detect_brands(a["text"], patterns)  # [idx, mentions, rank]
    platforms = sorted({a["platform"] for a in current})
    prompts = list(dict.fromkeys(a["prompt"] for a in current))
    dates = sorted({a["date"] for a in current})

    def stats(answers: list[dict]) -> list[dict]:
        s = [{"answers": 0, "mentions": 0, "rank_sum": 0, "top": 0} for _ in brands]
        total = 0
        for a in answers:
            for b, count, rank in a["brands"]:
                s[b]["answers"] += 1
                s[b]["mentions"] += count
                s[b]["rank_sum"] += rank
                s[b]["top"] += rank == 1
                total += count
        for x in s:
            x["vis"] = x["answers"] / len(answers) if answers else 0
            x["sov"] = x["mentions"] / total if total else 0
            x["avg"] = round(x["rank_sum"] / x["answers"], 1) if x["answers"] else None
        return s

    by_platform = {p: [a for a in current if a["platform"] == p] for p in platforms}
    overall = stats(current)
    per_pf = {p: stats(v) for p, v in by_platform.items()}

    wb = Workbook()
    wb.remove(wb.active)

    # 1. Summary
    summary_rows = []
    for label, answers, s in [("All platforms", current, overall)] + [(platform_name(p), by_platform[p], per_pf[p]) for p in platforms]:
        o = s[own]
        summary_rows.append([label, len(answers), o["answers"], o["vis"], o["avg"], o["top"], o["sov"]])
    write_sheet(
        wb, "Summary",
        ["Platform", "AI answers", f"Answers naming {own_brand}", "Visibility", "Avg. position", "Ranked #1", "Share of voice"],
        summary_rows, widths={1: 22, 3: 28}, pct_cols=(4, 7),
        intro=[
            f"AI visibility report: {project['name']}",
            f"Your brand: {own_brand}",
            f"Prompts: {len(prompts)}   Platforms: {', '.join(platform_name(p) for p in platforms)}   Scrape date(s): {', '.join(dates)}",
            "Visibility = share of answers that name the brand. Position = order of first mention among tracked brands (1 = named first).",
            "Share of voice = the brand's mentions as a share of all tracked-brand mentions.",
        ],
    )

    # 2. Brand leaderboard
    lb = []
    for i in sorted(range(len(brands)), key=lambda i: (-overall[i]["answers"], -overall[i]["top"])):
        o = overall[i]
        if not o["answers"] and i != own:
            continue
        lb.append([brands[i], o["answers"], o["vis"], o["sov"], o["avg"], o["top"], o["mentions"]]
                  + [per_pf[p][i]["vis"] for p in platforms])
    write_sheet(wb, "Brand leaderboard",
                ["Brand", "Answers naming it", "Visibility", "Share of voice", "Avg. position", "Ranked #1", "Mentions"]
                + [f"Visibility – {platform_name(p)}" for p in platforms],
                lb, widths={1: 24}, pct_cols=(3, 4, *range(8, 8 + len(platforms))))

    # 3. Brand websites cited
    site = defaultdict(lambda: {"answers": 0, "prompts": set(), "pages": set(), "cites": 0, "pf": Counter()})
    for a in current:
        seen = set()
        for u in a["cited"]:
            b = owner(u)
            if not b:
                continue
            e = site[b]
            e["cites"] += 1
            e["pages"].add(u)
            e["prompts"].add(a["prompt"])
            if b not in seen:
                e["answers"] += 1
                e["pf"][a["platform"]] += 1
                seen.add(b)
    site_rows = [[b, e["answers"], e["answers"] / len(current), len(e["prompts"]), len(e["pages"]), e["cites"]]
                 + [e["pf"][p] for p in platforms]
                 for b, e in sorted(site.items(), key=lambda kv: -kv[1]["answers"])]
    if own_brand not in site:
        site_rows.append([own_brand, 0, 0, 0, 0, 0] + [0 for _ in platforms])
    write_sheet(wb, "Brand websites cited",
                ["Brand website", "Answers citing it", "Share of answers", "Prompts", "Pages cited", "Citations"]
                + [f"Answers – {platform_name(p)}" for p in platforms],
                site_rows, widths={1: 24}, pct_cols=(3,))

    # 4. Top domains (current run)
    dom = defaultdict(lambda: {"answers": 0, "prompts": set(), "pages": set(), "read_only": 0, "pf": Counter()})
    for a in current:
        seen = set()
        for u in a["cited"]:
            d = data.domain_of(u)
            e = dom[d]
            e["pages"].add(u)
            e["prompts"].add(a["prompt"])
            if d not in seen:
                e["answers"] += 1
                e["pf"][a["platform"]] += 1
                seen.add(d)
        for d in {data.domain_of(u) for u in a["read_only"]}:
            dom[d]["read_only"] += 1
    dom_rows = [[d, owner("https://" + d) or "Third-party", url_type("https://" + d), e["answers"], e["answers"] / len(current),
                 len(e["prompts"]), len(e["pages"]), e["read_only"]] + [e["pf"][p] for p in platforms]
                for d, e in sorted(dom.items(), key=lambda kv: (-kv[1]["answers"], -kv[1]["read_only"]))]
    write_sheet(wb, "Top domains",
                ["Domain", "Owner", "Type", "Answers citing it", "Share of answers", "Prompts", "Pages cited",
                 "Answers that read it without citing"] + [f"Answers – {platform_name(p)}" for p in platforms],
                dom_rows, widths={1: 34, 2: 22, 8: 26}, pct_cols=(5,))

    # 5. Prompt summary
    ps_rows = []
    for a in sorted(current, key=lambda a: (prompts.index(a["prompt"]), a["platform"])):
        ordered = sorted(a["brands"], key=lambda x: x[2])
        own_hit = next((x for x in a["brands"] if x[0] == own), None)
        ps_rows.append([
            a["prompt"], platform_name(a["platform"]),
            "Yes" if own_hit else "No", own_hit[2] if own_hit else None,
            brands[ordered[0][0]] if ordered else "",
            ", ".join(f"{x[2]}. {brands[x[0]]}" for x in ordered),
            len(a["cited"]), len(a["read"]) or None, len(a["queries"]) or None,
        ])
    write_sheet(wb, "Prompt summary",
                ["Prompt", "Platform", f"{own_brand} named", f"{own_brand} position", "Named first",
                 "Brands in order of first mention", "Sources cited", "Pages read", "Searches run (ChatGPT)"],
                ps_rows, widths={1: 48, 3: 20, 4: 20, 5: 22, 6: 70, 8: 20, 9: 20}, wrap_cols=(1, 6))

    # 6/7. Sources from every platform in the current run, plus older ChatGPT runs
    runs = [("Current run", current)]
    for spec in args.previous:
        label, path = spec.split("=", 1)
        runs.append((label, read_run(data.REPO / path, {"chatgpt"})))
    src_rows = []
    uniq = defaultdict(lambda: {"cited": 0, "read": 0, "prompts": set(), "runs": set(), "platforms": set()})
    for label, answers in runs:
        for a in answers:
            for status, urls in (("Cited", a["cited"]), ("Read, not cited", a["read_only"])):
                for u in urls:
                    src_rows.append([label, platform_name(a["platform"]), a["prompt"], status, data.domain_of(u),
                                     owner(u) or "Third-party", url_type(u), u])
                    e = uniq[u]
                    e["cited" if status == "Cited" else "read"] += 1
                    e["prompts"].add(a["prompt"])
                    e["runs"].add(label)
                    e["platforms"].add(platform_name(a["platform"]))
    intro = [
        "Every source each AI cited or read for each prompt.",
        "Cited = linked in the answer. Read, not cited = returned by the AI's web search but not linked in the answer.",
    ]
    if args.previous:
        intro.append("Earlier ChatGPT runs asked every prompt as a follow-up in one shared chat, so their answers were "
                     "influenced by earlier prompts; they are included here for their sources only.")
    write_sheet(wb, "Sources (per prompt)",
                ["Run", "Platform", "Prompt", "Status", "Domain", "Owner", "Type", "URL"],
                src_rows, widths={1: 22, 2: 12, 3: 44, 4: 16, 5: 30, 6: 20, 7: 18, 8: 90}, intro=intro)
    uniq_rows = [[data.domain_of(u), owner(u) or "Third-party", url_type(u), e["cited"], e["read"],
                  len(e["prompts"]), ", ".join(sorted(e["platforms"])), ", ".join(sorted(e["runs"])), u]
                 for u, e in sorted(uniq.items(), key=lambda kv: (-kv[1]["cited"], -kv[1]["read"], kv[0]))]
    write_sheet(wb, "Source list (unique)",
                ["Domain", "Owner", "Type", "Times cited", "Times read, not cited", "Prompts", "Platforms", "Runs", "URL"],
                uniq_rows, widths={1: 30, 2: 20, 3: 18, 7: 18, 8: 30, 9: 90})

    # 8. Raw data
    raw_rows = []
    for label, answers in [("Current run", current)] + runs[1:]:
        for a in sorted(answers, key=lambda a: (a["prompt"], a["platform"])):
            raw_rows.append([label, platform_name(a["platform"]), a["date"], a["prompt"], a["text"],
                             "\n".join(a["cited"]), "\n".join(a["queries"]), "\n".join(a["read_only"])])
    write_sheet(wb, "Raw data",
                ["Run", "Platform", "Date", "Prompt", "Answer", "Cited URLs", "Searches run (ChatGPT)",
                 "Pages read, not cited"],
                raw_rows, widths={1: 22, 2: 12, 3: 12, 4: 40, 5: 100, 6: 60, 7: 45, 8: 60}, wrap_cols=(4, 5, 6, 7, 8))

    out = Path(args.out)
    wb.save(out)
    print(f"wrote {out} ({len(current)} answers, {len(src_rows)} source rows, {len(uniq_rows)} unique URLs)")


if __name__ == "__main__":
    main()
